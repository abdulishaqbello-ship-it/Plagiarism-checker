from flask import Flask, request, send_file
from docx import Document
from pypdf import PdfReader
from difflib import SequenceMatcher
from io import BytesIO
from datetime import datetime
from html import escape
import re
import json
import os

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024

# ---------------- PROJECT DETAILS ----------------
SCHOOL_NAME = "Federal University Dutsin-Ma (FUDMA)"
GROUP_NUMBER = "13"
PROJECT_TITLE = "Plagiarism Checker"
SUPERVISOR = "Engr. Nura"

# Edit these placeholders to add your group members.
GROUP_MEMBERS = [
    {"name": "ABDULRAHMAN ISHAQ BELLO", "matric": "Enter matric number", "department": "Enter department"},
    {"name": "Enter member 2 name", "matric": "Enter matric number", "department": "Enter department"},
    {"name": "Enter member 3 name", "matric": "Enter matric number", "department": "Enter department"},
    {"name": "Enter member 4 name", "matric": "Enter matric number", "department": "Enter department"},
]

HISTORY_FILE = "comparison_history.json"
comparison_history = []


def load_history():
    global comparison_history
    try:
        if os.path.exists(HISTORY_FILE):
            with open(HISTORY_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                comparison_history = data if isinstance(data, list) else []
    except (OSError, json.JSONDecodeError):
        comparison_history = []


def save_history():
    try:
        with open(HISTORY_FILE, "w", encoding="utf-8") as f:
            json.dump(comparison_history[:20], f, ensure_ascii=False, indent=2)
    except OSError:
        pass


def extract_text(file):
    name = (file.filename or "").lower()
    data = file.read()

    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)

    if name.endswith(".pdf"):
        pdf = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in pdf.pages)

    if name.endswith(".txt"):
        return data.decode("utf-8-sig", errors="replace")

    raise ValueError("Please upload a Word (.docx), PDF (.pdf), or text (.txt) file.")


def normalize_text(text):
    return re.sub(r"\s+", " ", text.casefold()).strip()


def matching_sentences(text1, text2):
    sentences1 = re.split(r"(?<=[.!?])\s+|\n+", text1)
    sentences2 = [normalize_text(s) for s in re.split(r"(?<=[.!?])\s+|\n+", text2)]
    targets = {s for s in sentences2 if len(s) >= 12}
    matches, seen = [], set()
    for sentence in sentences1:
        cleaned = normalize_text(sentence)
        if len(cleaned) >= 12 and cleaned in targets and cleaned not in seen:
            matches.append(sentence.strip())
            seen.add(cleaned)
    return matches


def status_for_score(score):
    if score < 30:
        return "Low similarity"
    if score < 60:
        return "Moderate similarity"
    return "High similarity"


def page_template(title, body, active="home"):
    nav_items = [
        ("home", "/", "⌂  Home"),
        ("about", "/about", "ⓘ  About Project"),
        ("history", "/history", "◷  History"),
    ]
    nav = "".join(
        f'<a class="nav-link {"active" if key == active else ""}" href="{href}">{label}</a>'
        for key, href, label in nav_items
    )
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="theme-color" content="#123b2a">
<title>{escape(title)} | {escape(SCHOOL_NAME)}</title>
<style>
:root {{
 --green:#176b45; --deep:#103b2b; --mint:#dff4e8; --ink:#172b22;
 --muted:#607368; --paper:#ffffff; --line:#e1ebe4; --bg:#f3f7f4;
 --gold:#f0bd59; --shadow:0 12px 35px rgba(16,59,43,.08);
}}
* {{ box-sizing:border-box; }}
body {{ margin:0; font-family:system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",Arial,sans-serif; background:var(--bg); color:var(--ink); line-height:1.6; }}
.topline {{ height:5px; background:linear-gradient(90deg,var(--gold),#78c69b,var(--green)); }}
header {{ color:white; background:radial-gradient(circle at top right,#287b55 0,transparent 38%),linear-gradient(135deg,#103b2b,#176b45); padding:26px 16px 30px; text-align:center; }}
.brand {{ display:flex; align-items:center; justify-content:center; gap:12px; max-width:850px; margin:0 auto 13px; }}
.brand-mark {{ display:flex; align-items:center; justify-content:center; width:54px; height:54px; border:1px solid #ffffff70; border-radius:16px; background:#ffffff15; font-size:25px; font-weight:800; }}
.brand-text {{ text-align:left; }}
.brand-text .eyebrow {{ margin:0; font-size:11px; letter-spacing:2px; text-transform:uppercase; color:#c7e8d5; }}
.brand-text h1 {{ margin:0; font-size:clamp(17px,4.3vw,24px); line-height:1.25; }}
header .project-name {{ margin:13px 0 3px; font-size:clamp(26px,6vw,38px); line-height:1.15; letter-spacing:-.7px; }}
header .subtitle {{ margin:8px auto 0; max-width:560px; color:#dcefe4; font-size:14px; }}
.badges {{ display:flex; justify-content:center; flex-wrap:wrap; gap:8px; margin-top:17px; }}
.badge {{ padding:5px 11px; border:1px solid #ffffff45; border-radius:30px; color:#fff; background:#ffffff12; font-size:12px; }}
nav {{ display:flex; justify-content:center; flex-wrap:wrap; gap:8px; padding:12px; background:#fff; border-bottom:1px solid var(--line); }}
.nav-link {{ text-decoration:none; padding:9px 13px; color:#4c6255; font-weight:650; font-size:13px; border-radius:9px; }}
.nav-link:hover,.nav-link.active {{ background:var(--mint); color:var(--green); }}
main {{ max-width:900px; margin:24px auto; padding:0 14px 30px; }}
.card {{ background:var(--paper); padding:clamp(17px,4vw,27px); margin:16px 0; border:1px solid var(--line); border-radius:17px; box-shadow:var(--shadow); }}
.card-heading {{ display:flex; align-items:flex-start; gap:12px; margin-bottom:8px; }}
.icon {{ flex:0 0 42px; width:42px; height:42px; display:flex; align-items:center; justify-content:center; border-radius:12px; background:var(--mint); color:var(--green); font-size:21px; }}
h2 {{ margin:0; color:var(--deep); font-size:22px; line-height:1.3; }}
h3 {{ color:var(--deep); font-size:17px; margin:22px 0 10px; }}
.muted {{ color:var(--muted); font-size:14px; margin:6px 0 0; }}
.upload-grid {{ display:grid; grid-template-columns:1fr 1fr; gap:14px; margin:22px 0; }}
.upload-tile {{ padding:17px; background:#f8fbf9; border:1.5px dashed #b9d2c2; border-radius:13px; min-width:0; }}
.upload-tile label {{ display:block; font-weight:750; color:var(--deep); margin-bottom:9px; }}
.upload-icon {{ font-size:22px; margin-bottom:7px; }}
input[type=file] {{ display:block; width:100%; max-width:100%; font-size:13px; color:#4c6255; }}
input[type=file]::file-selector-button {{ padding:8px 10px; margin-right:8px; border:0; border-radius:7px; background:#e0eee5; color:var(--deep); font-weight:650; }}
.button {{ display:inline-flex; justify-content:center; align-items:center; gap:8px; padding:12px 17px; border:0; border-radius:10px; background:var(--green); color:white; text-decoration:none; font-size:14px; font-weight:750; cursor:pointer; box-shadow:0 5px 12px #176b4522; }}
.button:hover {{ background:var(--deep); }}
.button.secondary {{ background:#e6f2ea; color:var(--green); box-shadow:none; }}
.button-row {{ display:flex; flex-wrap:wrap; gap:10px; align-items:center; }}
.full {{ width:100%; }}
.result-card {{ border-top:4px solid var(--green); }}
.score-row {{ display:flex; align-items:center; justify-content:space-between; flex-wrap:wrap; gap:12px; margin:17px 0 8px; }}
.score {{ font-size:clamp(32px,8vw,46px); font-weight:850; color:var(--green); letter-spacing:-1px; line-height:1.1; }}
.status {{ display:inline-block; padding:6px 11px; border-radius:30px; background:#e4f4e8; color:#176b45; font-size:12px; font-weight:800; }}
progress {{ display:block; width:100%; height:13px; margin:14px 0 18px; border:0; border-radius:99px; overflow:hidden; accent-color:var(--green); }}
progress::-webkit-progress-bar {{ background:#e5eee8; border-radius:99px; }}
progress::-webkit-progress-value {{ background:linear-gradient(90deg,#6ebc8b,var(--green)); border-radius:99px; }}
.notice {{ background:#fff7e5; border:1px solid #f3dfad; color:#6c5420; padding:13px 15px; border-radius:11px; font-size:13px; }}
.error {{ background:#fff0f0; border:1px solid #f1c6c6; color:#8d1010; padding:13px 15px; border-radius:11px; }}
.match {{ background:#f0faf3; border-left:4px solid #3c9a62; padding:12px 14px; margin:9px 0; border-radius:0 9px 9px 0; overflow-wrap:anywhere; }}
.empty {{ padding:14px; background:#f7faf8; border-radius:10px; color:var(--muted); }}
table {{ width:100%; border-collapse:collapse; font-size:13px; }}
.table-wrap {{ overflow-x:auto; }}
th,td {{ text-align:left; padding:11px 9px; border-bottom:1px solid var(--line); overflow-wrap:anywhere; }}
th {{ color:var(--deep); background:#f4f8f5; }}
.small {{ font-size:12px; color:var(--muted); }}
footer {{ text-align:center; padding:22px 14px; color:var(--muted); font-size:12px; }}
.footer-line {{ width:48px; height:3px; border-radius:4px; background:var(--gold); margin:0 auto 10px; }}
@media(max-width:580px) {{
 .upload-grid {{ grid-template-columns:1fr; }}
 .brand {{ gap:9px; }}
 .brand-mark {{ width:46px; height:46px; border-radius:13px; }}
 .brand-text .eyebrow {{ font-size:9px; letter-spacing:1.2px; }}
 .nav-link {{ padding:8px 10px; font-size:12px; }}
 .button-row .button {{ flex:1 1 100%; }}
}}
</style>
</head>
<body>
<div class="topline"></div>
<header>
 <div class="brand">
  <div class="brand-mark">F</div>
  <div class="brand-text">
   <p class="eyebrow">University Project · Group {escape(GROUP_NUMBER)}</p>
   <h1>{escape(SCHOOL_NAME)}</h1>
  </div>
 </div>
 <div class="project-name">{escape(PROJECT_TITLE)}</div>
 <p class="subtitle">A simple, user-friendly tool for comparing text similarity across documents.</p>
 <div class="badges"><span class="badge">PDF supported</span><span class="badge">Word supported</span><span class="badge">Text files supported</span></div>
</header>
<nav>{nav}</nav>
<main>{body}</main>
<footer><div class="footer-line"></div><strong>{escape(SCHOOL_NAME)}</strong><br>Group {escape(GROUP_NUMBER)} · Supervised by {escape(SUPERVISOR)}<br><span class="small">Similarity results are estimates and are not proof of plagiarism.</span></footer>
</body>
</html>"""


@app.route("/", methods=["GET", "POST"])
def home():
    result_html = ""
    error_html = ""

    if request.method == "POST":
        try:
            file1 = request.files.get("document1")
            file2 = request.files.get("document2")
            if not file1 or not file2 or not file1.filename or not file2.filename:
                raise ValueError("Please choose both documents.")

            filename1 = os.path.basename(file1.filename)
            filename2 = os.path.basename(file2.filename)
            allowed = (".docx", ".pdf", ".txt")
            if not filename1.lower().endswith(allowed) or not filename2.lower().endswith(allowed):
                raise ValueError("Supported formats are Word (.docx), PDF (.pdf), and text (.txt).")

            raw1 = extract_text(file1)
            raw2 = extract_text(file2)
            text1 = normalize_text(raw1)
            text2 = normalize_text(raw2)
            if not text1 or not text2:
                raise ValueError("No readable text found. Scanned/image-only PDFs need OCR.")

            score = SequenceMatcher(None, text1, text2, autojunk=False).ratio() * 100
            status = status_for_score(score)
            matches = matching_sentences(raw1, raw2)

            item = {
                "date": datetime.now().strftime("%Y-%m-%d %H:%M"),
                "document1": filename1,
                "document2": filename2,
                "score": round(score, 2),
                "status": status,
            }
            comparison_history.insert(0, item)
            del comparison_history[20:]
            save_history()

            match_html = "".join(f'<div class="match">✓ &nbsp;{escape(s)}</div>' for s in matches)
            if not match_html:
                match_html = '<div class="empty">No exact matching sentences of sufficient length were found. Other similar wording may still exist.</div>'

            result_html = f"""
            <section class="card result-card" id="result">
              <div class="card-heading"><div class="icon">📊</div><div><h2>Comparison Result</h2><p class="muted">Your document analysis is complete.</p></div></div>
              <p><strong>First document:</strong> {escape(filename1)}</p>
              <p><strong>Second document:</strong> {escape(filename2)}</p>
              <div class="score-row"><div class="score">{score:.2f}%</div><span class="status">{escape(status)}</span></div>
              <progress max="100" value="{score:.2f}">{score:.2f}%</progress>
              <div class="notice"><strong>Please note:</strong> This is a basic text-similarity estimate. It does not prove plagiarism or search the internet/external databases.</div>
              <h3>Exact matching sentences</h3>
              {match_html}
              <div class="button-row" style="margin-top:18px">
                <form method="POST" action="/download-report">
                  <input type="hidden" name="document1" value="{escape(filename1, quote=True)}">
                  <input type="hidden" name="document2" value="{escape(filename2, quote=True)}">
                  <input type="hidden" name="score" value="{score:.2f}">
                  <input type="hidden" name="status" value="{escape(status, quote=True)}">
                  <input type="hidden" name="matches" value="{escape(json.dumps(matches, ensure_ascii=False), quote=True)}">
                  <button class="button" type="submit">⬇ Download Report</button>
                </form>
                <a class="button secondary" href="/">＋ New Comparison</a>
              </div>
            </section>"""
        except Exception as exc:
            error_html = f'<p class="error"><strong>Unable to compare:</strong> {escape(str(exc))}</p>'

    body = f"""
    <section class="card">
      <div class="card-heading"><div class="icon">⇄</div><div><h2>Compare Two Documents</h2><p class="muted">Select two files to calculate how similar their text is.</p></div></div>
      <form method="POST" enctype="multipart/form-data">
        <div class="upload-grid">
          <div class="upload-tile">
            <div class="upload-icon">📄</div>
            <label for="document1">Document 1</label>
            <input id="document1" type="file" name="document1" accept=".docx,.pdf,.txt" required>
            <p class="small">PDF, Word (.docx), or text (.txt)</p>
          </div>
          <div class="upload-tile">
            <div class="upload-icon">📑</div>
            <label for="document2">Document 2</label>
            <input id="document2" type="file" name="document2" accept=".docx,.pdf,.txt" required>
            <p class="small">PDF, Word (.docx), or text (.txt)</p>
          </div>
        </div>
        <button class="button full" type="submit">Compare Documents &nbsp; →</button>
      </form>
      {error_html}
    </section>
    {result_html}
    <section class="card">
      <div class="card-heading"><div class="icon">✦</div><div><h2>How it works</h2><p class="muted">A quick guide to your document checker.</p></div></div>
      <p><strong>1.</strong> Upload two supported documents.</p>
      <p><strong>2.</strong> Select <em>Compare Documents</em>.</p>
      <p><strong>3.</strong> Review the estimated similarity score and matching sentences.</p>
      <p class="small">Note: scanned PDFs containing only images may need OCR before text can be extracted.</p>
    </section>
    """
    return page_template("Plagiarism Checker", body, "home")


@app.route("/about")
def about():
    rows = "".join(
        f"<tr><td>{i}</td><td>{escape(m.get('name', ''))}</td><td>{escape(m.get('matric', ''))}</td><td>{escape(m.get('department', ''))}</td></tr>"
        for i, m in enumerate(GROUP_MEMBERS, start=1)
    )
    body = f"""
    <section class="card">
      <div class="card-heading"><div class="icon">ⓘ</div><div><h2>About Our Project</h2><p class="muted">Project and group information.</p></div></div>
      <div class="table-wrap"><table><tbody>
        <tr><th>University</th><td>{escape(SCHOOL_NAME)}</td></tr>
        <tr><th>Group Number</th><td>{escape(GROUP_NUMBER)}</td></tr>
        <tr><th>Project Title</th><td>{escape(PROJECT_TITLE)}</td></tr>
        <tr><th>Supervisor</th><td>{escape(SUPERVISOR)}</td></tr>
      </tbody></table></div>
      <h3>Group Members</h3>
      <p class="small">Edit the GROUP_MEMBERS list near the top of app.py to add the correct names, matric numbers, and departments.</p>
      <div class="table-wrap"><table><thead><tr><th>No.</th><th>Full name</th><th>Matric number</th><th>Department</th></tr></thead><tbody>{rows}</tbody></table></div>
    </section>"""
    return page_template("About Our Project", body, "about")


@app.route("/history")
def history():
    if comparison_history:
        rows = "".join(
            f"<tr><td>{escape(str(x.get('date','')))}</td><td>{escape(str(x.get('document1','')))}</td><td>{escape(str(x.get('document2','')))}</td><td><strong>{escape(str(x.get('score','')))}%</strong></td><td>{escape(str(x.get('status','')))}</td></tr>"
            for x in comparison_history
        )
        table = f'<div class="table-wrap"><table><thead><tr><th>Date</th><th>Document 1</th><th>Document 2</th><th>Similarity</th><th>Status</th></tr></thead><tbody>{rows}</tbody></table></div>'
    else:
        table = '<div class="empty">No comparisons have been recorded yet. Your future comparisons will appear here.</div>'
    body = f"""
    <section class="card">
      <div class="card-heading"><div class="icon">◷</div><div><h2>Comparison History</h2><p class="muted">Your 20 most recent comparisons on this device.</p></div></div>
      {table}
      <p class="small">History is stored locally in comparison_history.json.</p>
    </section>"""
    return page_template("Comparison History", body, "history")


@app.route("/download-report", methods=["POST"])
def download_report():
    document1 = request.form.get("document1", "Document 1")
    document2 = request.form.get("document2", "Document 2")
    score = request.form.get("score", "N/A")
    status = request.form.get("status", "N/A")
    try:
        matches = json.loads(request.form.get("matches", "[]"))
        if not isinstance(matches, list):
            matches = []
    except json.JSONDecodeError:
        matches = []

    report = [
        SCHOOL_NAME, f"Project: {PROJECT_TITLE}", f"Group Number: {GROUP_NUMBER}",
        f"Supervisor: {SUPERVISOR}", "", "DOCUMENT SIMILARITY REPORT",
        f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"First document: {document1}", f"Second document: {document2}",
        f"Similarity: {score}%", f"Status: {status}", "", "Exact matching sentences:",
    ]
    report.extend(f"- {s}" for s in matches)
    report.extend(["", "Note: This is a basic text-similarity estimate, not proof of plagiarism.",
                   "The application does not search the internet or external plagiarism databases."])
    output = BytesIO("\n".join(report).encode("utf-8"))
    output.seek(0)
    return send_file(output, mimetype="text/plain; charset=utf-8", as_attachment=True,
                     download_name="plagiarism_comparison_report.txt")


@app.errorhandler(413)
def too_large(_error):
    return page_template("File Too Large",
        '<section class="card"><h2>File too large</h2><p class="error">Please keep the combined upload size below 20 MB.</p><a class="button" href="/">Go back</a></section>'), 413


load_history()

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
